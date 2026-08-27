from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook

from src.models import DataQuality, OpportunityCandidate
from src.reporting.export import export_scan_results, results_to_frame


def candidate() -> OpportunityCandidate:
    return OpportunityCandidate(
        ticker="TEST",
        current_price=80.0,
        price_basis="adjusted_close",
        observation_date=date(2026, 1, 1),
        drawdown_52w=-0.2,
        decline_severity_score=30.0,
        is_candidate=True,
        candidate_reasons=["drawdown_52w=-20.00% <= -20.00%"],
        data_quality=DataQuality.MEDIUM,
        metric_statuses={},
        missing_metrics=[],
        sources=[{"name": "synthetic_test_fixture"}],
        retrieved_at=datetime(2026, 1, 1, tzinfo=UTC),
    )


def test_results_frame_serializes_nested_audit_fields() -> None:
    frame = results_to_frame([candidate()])
    assert frame.loc[0, "ticker"] == "TEST"
    assert "synthetic_test_fixture" in frame.loc[0, "sources"]


def test_exports_csv_and_workbook_sheets(tmp_path: Path) -> None:
    paths = export_scan_results([candidate()], tmp_path)
    assert {path.suffix for path in paths} == {".csv", ".xlsx"}
    workbook_path = next(path for path in paths if path.suffix == ".xlsx")
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook.sheetnames == ["Candidates", "All Results", "Run Metadata"]

