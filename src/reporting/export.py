"""CSV and styled Excel exports for integrated scan results."""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.models import OpportunityCandidate

EXPORT_COLUMNS = [
    "rank",
    "ticker",
    "company",
    "country",
    "sector",
    "exchange",
    "current_price",
    "currency",
    "current_price_usd",
    "current_price_eur",
    "market_cap",
    "market_cap_currency",
    "market_cap_usd",
    "market_cap_eur",
    "fx_as_of",
    "fx_metrics",
    "observation_date",
    "price_basis",
    "drawdown_ath",
    "drawdown_52w",
    "drawdown_1m",
    "drawdown_3m",
    "drawdown_6m",
    "drawdown_ytd",
    "drawdown_1y",
    "return_1m",
    "return_3m",
    "return_6m",
    "return_ytd",
    "return_1y",
    "distance_ma50",
    "distance_ma200",
    "rsi",
    "volatility",
    "beta",
    "relative_benchmark_performance",
    "relative_sector_performance",
    "fundamental_quality_score",
    "fundamental_quality_observed_score",
    "fundamental_quality_coverage",
    "fundamental_growth_score",
    "fundamental_profitability_score",
    "fundamental_balance_sheet_score",
    "fundamental_cash_flow_score",
    "fundamental_data_quality",
    "fundamental_status",
    "fundamental_as_of",
    "fundamental_latest_period_end",
    "fundamental_metrics",
    "valuation_score",
    "valuation_observed_score",
    "valuation_coverage",
    "valuation_status",
    "valuation_data_quality",
    "valuation_as_of",
    "valuation_price",
    "valuation_price_date",
    "pe_current",
    "pe_historical_median",
    "pe_sector_median",
    "ev_ebitda_current",
    "ev_ebitda_historical_median",
    "ev_ebitda_sector_median",
    "price_fcf_current",
    "price_fcf_historical_median",
    "price_fcf_sector_median",
    "dcf_bear_value_per_share",
    "dcf_base_value_per_share",
    "dcf_bull_value_per_share",
    "normalized_value_per_share",
    "reverse_dcf_implied_revenue_growth",
    "valuation_metrics",
    "shock_category",
    "shock_nature",
    "temporary_shock_score",
    "temporary_shock_observed_score",
    "temporary_shock_coverage",
    "shock_status",
    "shock_data_quality",
    "shock_as_of",
    "shock_evidence_count",
    "shock_independent_source_count",
    "shock_specification_criteria_coverage",
    "shock_missing_criteria",
    "shock_conclusion",
    "shock_metrics",
    "historical_status",
    "historical_data_quality",
    "historical_as_of",
    "historical_analogue_count",
    "best_historical_similarity",
    "historical_metrics",
    "scenario_status",
    "scenario_data_quality",
    "scenario_as_of",
    "fair_value",
    "normalized_fair_value_scenario",
    "bear_target",
    "base_target",
    "bull_target",
    "tp1",
    "tp2",
    "tp3",
    "upside_base",
    "upside_bull",
    "downside_bear",
    "risk_reward",
    "fundamental_invalidation",
    "scenario_metrics",
    "normalization_score",
    "normalization_score_coverage",
    "catalyst_score",
    "catalyst_score_coverage",
    "future_growth_score",
    "risk_score",
    "risk_score_coverage",
    "opportunity_score",
    "opportunity_observed_score",
    "opportunity_score_coverage",
    "opportunity_score_status",
    "opportunity_data_quality",
    "confidence_score",
    "score_band",
    "scoring_metrics",
    "decline_severity_score",
    "is_candidate",
    "candidate_reasons",
    "data_quality",
    "metric_statuses",
    "missing_metrics",
    "sources",
    "retrieved_at",
]

PERCENT_COLUMNS = {
    "drawdown_ath",
    "drawdown_52w",
    "drawdown_1m",
    "drawdown_3m",
    "drawdown_6m",
    "drawdown_ytd",
    "drawdown_1y",
    "return_1m",
    "return_3m",
    "return_6m",
    "return_ytd",
    "return_1y",
    "distance_ma50",
    "distance_ma200",
    "volatility",
    "relative_benchmark_performance",
    "relative_sector_performance",
    "reverse_dcf_implied_revenue_growth",
    "upside_base",
    "upside_bull",
    "downside_bear",
}


def results_to_frame(results: Iterable[OpportunityCandidate]) -> pd.DataFrame:
    rows = []
    for result in results:
        row = result.model_dump(mode="json")
        row["candidate_reasons"] = json.dumps(row["candidate_reasons"], ensure_ascii=False)
        row["fundamental_metrics"] = json.dumps(
            row["fundamental_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["valuation_metrics"] = json.dumps(
            row["valuation_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["fx_metrics"] = json.dumps(
            row["fx_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["shock_metrics"] = json.dumps(
            row["shock_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["historical_metrics"] = json.dumps(
            row["historical_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["fundamental_invalidation"] = json.dumps(
            row["fundamental_invalidation"], ensure_ascii=False
        )
        row["scenario_metrics"] = json.dumps(
            row["scenario_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["scoring_metrics"] = json.dumps(
            row["scoring_metrics"], ensure_ascii=False, sort_keys=True
        )
        row["shock_missing_criteria"] = json.dumps(
            row["shock_missing_criteria"], ensure_ascii=False
        )
        row["metric_statuses"] = json.dumps(
            row["metric_statuses"], ensure_ascii=False, sort_keys=True
        )
        row["missing_metrics"] = json.dumps(row["missing_metrics"], ensure_ascii=False)
        row["sources"] = json.dumps(row["sources"], ensure_ascii=False)
        rows.append(row)
    return pd.DataFrame(rows, columns=EXPORT_COLUMNS)


def _style_sheet(sheet: Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    headers = {cell.value: cell.column for cell in sheet[1]}
    for name in PERCENT_COLUMNS:
        if name in headers:
            for cell in sheet[get_column_letter(headers[name])][1:]:
                cell.number_format = "0.00%"
    for name in (
        "current_price",
        "current_price_usd",
        "current_price_eur",
        "market_cap",
        "market_cap_usd",
        "market_cap_eur",
        "valuation_price",
        "pe_current",
        "ev_ebitda_current",
        "price_fcf_current",
        "dcf_bear_value_per_share",
        "dcf_base_value_per_share",
        "dcf_bull_value_per_share",
        "normalized_value_per_share",
        "fair_value",
        "normalized_fair_value_scenario",
        "bear_target",
        "base_target",
        "bull_target",
        "tp1",
        "tp2",
        "tp3",
        "risk_reward",
        "valuation_score",
        "temporary_shock_score",
        "normalization_score",
        "catalyst_score",
        "future_growth_score",
        "risk_score",
        "opportunity_score",
        "opportunity_observed_score",
        "confidence_score",
        "best_historical_similarity",
        "beta",
        "decline_severity_score",
        "rsi",
    ):
        if name in headers:
            for cell in sheet[get_column_letter(headers[name])][1:]:
                cell.number_format = "0.00"

    for column_cells in sheet.columns:
        header = str(column_cells[0].value or "")
        sample_width = max((len(str(cell.value or "")) for cell in column_cells[:100]), default=0)
        wide_columns = {
            "candidate_reasons",
            "fundamental_metrics",
            "metric_statuses",
            "shock_conclusion",
            "shock_metrics",
            "historical_metrics",
            "fundamental_invalidation",
            "scenario_metrics",
            "scoring_metrics",
            "shock_missing_criteria",
            "sources",
            "valuation_metrics",
        }
        limit = 60 if header in wide_columns else 28
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(sample_width + 2, len(header) + 2, 10), limit
        )


def export_scan_results(
    results: list[OpportunityCandidate],
    output_dir: str | Path,
    *,
    filename_prefix: str = "stock_opportunity_scan",
    write_csv: bool = True,
    write_excel: bool = True,
    run_metadata: dict[str, object] | None = None,
) -> list[Path]:
    """Export all auditable rows and a candidate-only Excel view."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    base = destination / f"{filename_prefix}_{timestamp}"
    frame = results_to_frame(results)
    paths: list[Path] = []

    if write_csv:
        csv_path = base.with_suffix(".csv")
        frame.to_csv(csv_path, index=False)
        paths.append(csv_path)

    if write_excel:
        excel_path = base.with_suffix(".xlsx")
        candidates = frame[frame["is_candidate"] == True]  # noqa: E712 - pandas mask
        metadata = {
            "generated_at": datetime.now(UTC).isoformat(),
            "result_count": len(frame),
            "candidate_count": len(candidates),
            "note": (
                "All scores are coverage-adjusted analytical indicators out of "
                "100, not investment recommendations or return/normalization "
                "probabilities. Risk Score is resilience: higher means lower "
                "measured risk."
            ),
            **(run_metadata or {}),
        }
        metadata_frame = pd.DataFrame(
            [
                {
                    "key": key,
                    "value": (
                        json.dumps(value) if isinstance(value, (dict, list)) else value
                    ),
                }
                for key, value in metadata.items()
            ]
        )
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            candidates.to_excel(writer, sheet_name="Candidates", index=False)
            frame.to_excel(writer, sheet_name="All Results", index=False)
            metadata_frame.to_excel(writer, sheet_name="Run Metadata", index=False)
            for sheet in writer.book.worksheets:
                _style_sheet(sheet)
        paths.append(excel_path)

    return paths
